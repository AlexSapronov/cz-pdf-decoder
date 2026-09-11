"""Фиксация контракта вывода Excel.

Цель — явно закрепить желаемое представление «Полный DM» в Excel
и тот факт, что serial НЕ экспортируется.
"""
import os

from openpyxl import load_workbook

from czdecoder.excel import save_excel, row_to_cells

GS = "\x1d"


def test_full_dm_gs_removed_on_export():
    """«Полный DM» с GS-разделителем экспортируется БЕЗ \\x1d.

    \\x1d (0x1D, Group Separator) входит в ILLEGAL_CHARACTERS_RE openpyxl и
    полностью УДАЛЯЕТСЯ clean_for_excel(). Итог — слитные «GTIN»+«serial»
    без разделителя.

    Это ОТЛИЧИЕ от baseline: там GS схлопывался в пробел (clean_for_excel
    применялся на decode stage, где str.split() резал по живому \\x1d и
    склеивал через пробел). Теперь decode сохраняет \\x1d, а export его
    удаляет целиком. Зафиксировано как намеренное поведение.
    """
    row = {
        "full_dm": "0104640638345218" + GS + "21SERIALXYZ",
        "gtin": "04640638345218",
        "serial": "SERIALXYZ",  # internal, НЕ должен попасть в вывод
        "pn": "Товар",
        "qty": "1",
        "file_name": "a.pdf",
        "page_num": 1,
    }
    cells = row_to_cells(row)
    assert len(cells) == 6  # ровно 6 колонок, без serial
    # «Полный DM»: GS полностью удалён, управляющего символа нет.
    assert "\x1d" not in cells[0]
    assert cells[0] == "010464063834521821SERIALXYZ"


def test_serial_not_in_headers_or_cells(tmp_path):
    """Serial — внутреннее поле: нет в HEADERS и не попадает в файл Excel."""
    from czdecoder.excel import HEADERS
    assert "serial" not in [h.lower() for h in HEADERS]
    assert len(HEADERS) == 6

    rows = [{
        "full_dm": "010464063834521821SERIAL",
        "gtin": "04640638345218",
        "serial": "SERIAL",  # internal — игнорируется при экспорте
        "pn": "Товар",
        "qty": "2",
        "file_name": "b.pdf",
        "page_num": 1,
    }]
    out = os.path.join(str(tmp_path), "out.xlsx")
    save_excel(rows, out)

    wb = load_workbook(out)
    ws = wb["DataMatrix"]
    header = [c.value for c in ws[1]]
    assert header == ["Полный DM", "GTIN", "PN", "Кол-во", "Файл", "Страница"]
    # строка данных: 6 ячеек, no serial value (SERIAL) нигде
    data_row = [c.value for c in ws[2]]
    assert len(data_row) == 6
    assert "SERIAL" not in data_row
